#!/usr/bin/env python3
"""Send cold emails via Postmark API. Supports .xlsx, .csv, and .json data files."""

import json
import os
import sys
import time
from datetime import datetime, timezone

import requests
from dotenv import load_dotenv

load_dotenv()

POSTMARK_SERVER_TOKEN = os.getenv("POSTMARK_SERVER_TOKEN")
SENDER_EMAIL = os.getenv("POSTMARK_SENDER_EMAIL", "rain@paperfox.ai")
POSTMARK_API_URL = "https://api.postmarkapp.com/email"

# ---------- template ----------

def load_template(path="email-template-v5.md"):
    with open(path, "r") as f:
        lines = f.readlines()
    # First line is subject
    subject_line = lines[0].strip()
    subject_template = subject_line.replace("Subject: ", "", 1)
    # Rest is body (strip leading/trailing blank lines)
    body_lines = lines[1:]
    body_template = "".join(body_lines).strip()
    return subject_template, body_template


def render(template_str, conference_name, first_name, platform="EasyChair"):
    """Replace placeholders in template string."""
    result = template_str.replace("[Conference Name]", conference_name)
    if first_name == _FALLBACK_GREETING:
        result = result.replace("Hi [Name]", first_name)
    result = result.replace("[Name]", first_name)
    result = result.replace("[Platform]", platform)
    return result


def plain_to_html(text):
    """Convert plain text body to proper HTML email, preserving the Markdown link.

    Wraps content in <html><body>...</body></html> so that Postmark can inject
    its open-tracking pixel before the closing </body> tag.
    """
    import re
    # Convert markdown link [text](url) to <a> tag
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2">\1</a>', text)
    # Split into paragraphs
    paragraphs = text.split("\n\n")
    html_parts = []
    for p in paragraphs:
        p = p.strip()
        if p:
            # Replace single newlines with <br>
            p = p.replace("\n", "<br>\n")
            html_parts.append(f"<p>{p}</p>")
    body_content = "\n".join(html_parts)
    return f"<html><body>\n{body_content}\n</body></html>"


# ---------- name extraction ----------

import re as _re

_CREDENTIAL_RE = _re.compile(
    r',\s*(?:'
    r'S\.T\.|M\.T\.|M\.Eng\.?|M\.Kom\.?|M\.Sc\.?|M\.Cs\.?|M\.Med\.Ed\.?|'
    r'M\.Farm\.?|M\.Kes\.?|MNSc\.?|MM\.?|M\.InfoTech|M\.A\.D\.E|'
    r'S\.Si\.?|S\.Kom\.?|B\.Eng\.?|Ns\.?|'
    r'Ph\.?\s*D\.?|IEEE\s*Member|IEEE'
    r').*$', _re.IGNORECASE)

_TITLE_TOKENS = {
    "prof.", "prof", "professor", "profesor",
    "dr.", "dr",
    "assoc.", "associate",
    "ir.", "ts.", "apt.", "eng.",
    "mr.", "mrs.", "ms.",
    "dr.ing.", "dr.-ing.",
    "h.e.", "en.",
}

_FALLBACK_GREETING = "Dear Conference Chair"


def _normalize_case(word: str) -> str:
    parts = word.split("-")
    normalized = [p.capitalize() if p.isupper() and len(p) > 2 else p for p in parts]
    return "-".join(normalized)


def extract_greeting_name(full_name: str) -> str:
    """Extract the appropriate greeting name from a full name.

    Handles complex academic titles (Assoc. Prof. Ir. Dr.),
    credential suffixes (S.T., M.Eng., Ph.D.), and ALL-CAPS surnames.
    """
    name = (full_name or "").strip()
    if not name:
        return "there"

    name = _CREDENTIAL_RE.sub("", name).strip().rstrip(".")

    # Handle Prof.(Dr.) style
    name = _re.sub(r'Prof\.\s*\(Dr\.\)', 'Prof.', name, flags=_re.IGNORECASE)

    tokens = name.split()
    title_end = 0
    has_prof = False
    has_dr = False
    for i, tok in enumerate(tokens):
        low = tok.lower().rstrip(",")
        if low in _TITLE_TOKENS:
            title_end = i + 1
            if low in ("prof.", "prof", "professor", "profesor"):
                has_prof = True
            if low in ("dr.", "dr", "dr.ing.", "dr.-ing."):
                has_dr = True
        elif low.startswith("dr.") or low.startswith("dr-"):
            title_end = i + 1
            has_dr = True
        else:
            break

    remainder = tokens[title_end:]
    if not remainder:
        return _FALLBACK_GREETING

    remainder = [_normalize_case(w) for w in remainder]

    if has_prof:
        return f"Prof. {remainder[-1]}"
    if has_dr:
        return f"Dr. {remainder[-1]}"

    first = remainder[0]
    _NOT_NAMES = {"md", "st", "mt"}
    if len(first) <= 1 or first.endswith(".") or first.endswith(","):
        return _FALLBACK_GREETING
    if first.lower() in _NOT_NAMES:
        return _FALLBACK_GREETING

    return first


# ---------- data loaders ----------

_SKIP_EMAIL_PREFIXES = (
    "noreply@", "admin@", "webmaster@", "info@",
    "helpdesk@", "registrar@", "support@", "easychair",
)


def _deduplicate_and_filter(raw_rows):
    """Deduplicate by email, skip empty/bad emails and empty names.

    Applies the shared strict validator (masked, compound, placeholder, malformed)
    plus this script's role-account prefix filter (noreply/admin/info/…).
    """
    from database.crm_db import is_valid_email as _strict_is_valid_email

    recipients = []
    seen = set()
    for row in raw_rows:
        email = (row.get("chair_email") or "").strip()
        name = (row.get("chair_name") or "").strip()
        conf = (row.get("conference_short_name") or "").strip()
        if not email or not name:
            continue
        if email.lower() in seen:
            continue
        ok, reason = _strict_is_valid_email(email)
        if not ok:
            print(f"  [SKIP] invalid email '{email}' ({reason})")
            continue
        if any(email.lower().startswith(p) for p in _SKIP_EMAIL_PREFIXES):
            continue
        seen.add(email.lower())
        platform = (row.get("platform") or "EasyChair").strip()
        recipients.append({
            "chair_name": name,
            "first_name": extract_greeting_name(name),
            "chair_email": email,
            "conference_short_name": conf,
            "platform": platform,
        })
    return recipients


def load_xlsx(path):
    """Load recipients from an .xlsx file. Expects columns: Name, email, ..., Conference short name (col H)."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    raw = []
    for row in rows[1:]:
        raw.append({
            "chair_name": row[0] or "",
            "chair_email": row[1] or "",
            "conference_short_name": row[7] or "" if len(row) > 7 else "",
        })
    return _deduplicate_and_filter(raw)


def load_csv(path):
    """Load recipients from a .csv file. Expects columns: chair_name, chair_email, conference_short_name."""
    import csv
    with open(path, "r") as f:
        reader = csv.DictReader(f)
        raw = list(reader)
    return _deduplicate_and_filter(raw)


def load_json(path):
    """Load recipients from a .json file. Expects list of objects with chair_name, chair_email, conference_short_name."""
    with open(path, "r") as f:
        raw = json.load(f)
    return _deduplicate_and_filter(raw)


def load_recipients(path):
    """Auto-detect file type and load recipients."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return load_xlsx(path)
    elif ext == ".csv":
        return load_csv(path)
    elif ext == ".json":
        return load_json(path)
    else:
        print(f"ERROR: Unsupported file type '{ext}'. Use .xlsx, .csv, or .json")
        sys.exit(1)


# ---------- skip already sent ----------

def _get_already_sent_emails() -> set:
    """Query DB for all emails that were already sent (including bounced)."""
    try:
        from database.db_config import get_connection
        conn = get_connection()
        cur = conn.cursor()
        cur.execute("SELECT LOWER(contact_email) FROM emails")
        sent = {row[0] for row in cur.fetchall()}
        cur.close()
        conn.close()
        return sent
    except Exception:
        return set()


# ---------- sending ----------

def send_email(recipient, subject, plain_body, html_body, dry_run=True):
    """Send a single email via Postmark. Returns result dict.

    The rendered plain_body / html_body / platform are carried into the result
    so that save_to_db() can persist them on the emails row. This lets future
    follow-ups forward the exact content that was actually delivered.
    """
    result = {
        "email": recipient["chair_email"],
        "conference": recipient["conference_short_name"],
        "chair_name": recipient["chair_name"],
        "platform": recipient.get("platform"),
        "subject": subject,
        "plain_body": plain_body,
        "html_body": html_body,
    }

    if dry_run:
        result["status"] = "dry_run"
        return result

    try:
        # Use a Session with trust_env=False to bypass any proxy env vars that
        # may be set by .env (e.g. HTTP_PROXY=http://127.0.0.1:7892). Postmark
        # is publicly reachable; routing through a local proxy is never wanted.
        _session = requests.Session()
        _session.trust_env = False
        resp = _session.post(
            POSTMARK_API_URL,
            headers={
                "Accept": "application/json",
                "Content-Type": "application/json",
                "X-Postmark-Server-Token": POSTMARK_SERVER_TOKEN,
            },
            json={
                "From": SENDER_EMAIL,
                "To": recipient["chair_email"],
                "Subject": subject,
                "HtmlBody": html_body,
                "TextBody": plain_body,
                "MessageStream": "outbound",
                "TrackOpens": True,
                "TrackLinks": "HtmlOnly",
            },
            timeout=30,
        )
        data = resp.json()
        if resp.status_code == 200 and data.get("ErrorCode") == 0:
            result["status"] = "sent"
            result["postmark_message_id"] = data.get("MessageID", "")
        else:
            result["status"] = "error"
            result["error_code"] = data.get("ErrorCode")
            result["error_message"] = data.get("Message", "")
    except Exception as e:
        result["status"] = "error"
        result["error_message"] = str(e)

    result["sent_at"] = datetime.now(timezone.utc).isoformat()
    return result


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Send cold emails via Postmark")
    parser.add_argument("data_file", help="Path to recipient data file (.xlsx, .csv, or .json)")
    parser.add_argument("--send", action="store_true",
                        help="Actually send emails (without this flag, does a dry run)")
    parser.add_argument("--limit", type=int, default=0,
                        help="Limit number of emails to send (0 = no limit)")
    parser.add_argument("--template", default="email-template-v5.md",
                        help="Path to email template file")
    parser.add_argument("--test", action="store_true",
                        help="Use test database (crm_test) instead of production")
    args = parser.parse_args()

    if args.test:
        os.environ["USE_TEST_DB"] = "1"
        print("[TEST MODE] Using test database crm_test\n")

    if not POSTMARK_SERVER_TOKEN:
        print("ERROR: POSTMARK_SERVER_TOKEN not set in .env")
        sys.exit(1)

    # Load template
    subject_template, body_template = load_template(args.template)

    # Load recipients
    recipients = load_recipients(args.data_file)
    print(f"Loaded {len(recipients)} recipients from {args.data_file}")

    # Skip already-sent emails (query DB)
    already_sent = _get_already_sent_emails()
    if already_sent:
        before = len(recipients)
        recipients = [r for r in recipients if r["chair_email"].lower() not in already_sent]
        skipped = before - len(recipients)
        if skipped:
            print(f"Skipped {skipped} already-sent recipient(s)")

    if args.limit > 0:
        recipients = recipients[:args.limit]
        print(f"Limited to {len(recipients)} recipients")

    dry_run = not args.send
    if dry_run:
        print("\n" + "=" * 60)
        print("  DRY RUN — no emails will be sent")
        print("  Add --send flag to actually send")
        print("=" * 60)

    # Preview
    print(f"\nSender: {SENDER_EMAIL}")
    print(f"Recipients: {len(recipients)}")
    print()

    results = []
    for i, r in enumerate(recipients):
        platform = r.get("platform", "EasyChair")
        subject = render(subject_template, r["conference_short_name"], r["first_name"], platform)
        plain_body = render(body_template, r["conference_short_name"], r["first_name"], platform)
        html_body = plain_to_html(plain_body)

        # Show preview for first 3
        if i < 3 or dry_run:
            print(f"--- Email {i+1} ---")
            print(f"  To:      {r['chair_email']}")
            print(f"  Name:    {r['chair_name']}")
            print(f"  Subject: {subject}")
            if i < 2:  # Full preview for first 2
                print(f"  Body:\n{plain_body}")
            print()

        result = send_email(r, subject, plain_body, html_body, dry_run=dry_run)
        results.append(result)

        if not dry_run:
            status = result.get("status", "?")
            extra = ""
            if status == "error":
                extra = f"  ERR_CODE={result.get('error_code')!r}  ERR_MSG={result.get('error_message','')[:200]!r}"
            print(f"  [{i+1}/{len(recipients)}] {r['chair_email']:40s} {status}{extra}", flush=True)
            if i < len(recipients) - 1:
                time.sleep(0.5)

    # Summary
    sent = sum(1 for r in results if r["status"] == "sent")
    errors = sum(1 for r in results if r["status"] == "error")
    dry = sum(1 for r in results if r["status"] == "dry_run")

    print("=" * 60)
    if dry_run:
        print(f"DRY RUN complete: {dry} emails previewed")
        print("Run with --send to actually send")
    else:
        print(f"DONE: {sent} sent, {errors} errors")

    # Write results to database
    if not dry_run:
        save_to_db(results)


def save_to_db(results):
    """Write send results directly into PostgreSQL database."""
    from database.crm_db import upsert_contact, insert_email
    from database.db_config import get_connection

    conn = get_connection()
    added = 0
    for r in results:
        if r["status"] != "sent":
            continue
        upsert_contact(conn, r["email"], r.get("chair_name", ""), r.get("conference", ""),
                       source_platform=r.get("platform"))
        email_id = insert_email(
            conn, r["email"], r.get("postmark_message_id", ""),
            r.get("subject", ""), r.get("sent_at", ""),
            body_text=r.get("plain_body"),
            body_html=r.get("html_body"),
        )
        if email_id:
            added += 1
    conn.commit()
    conn.close()
    print(f"Database updated: {added} email(s) recorded")


if __name__ == "__main__":
    main()
